"""Import anonymous feedback from the first CSV/XLSX file in content-source/feedback."""

from __future__ import annotations

import csv
import json
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = ROOT / "content-source" / "feedback"
OUTPUT = ROOT / "src" / "app" / "data" / "feedbacks.generated.json"
PREFERRED_HEADERS = ("feedback", "comment", "testimonial", "review")
PLACEHOLDERS = {"n/a", "na", "nil", "none", "no", "nope", "-", ".", "nothing"}


def clean(value: object) -> str:
    return " ".join(str(value or "").strip().split())


def choose_columns(headers: list[str], rows: list[list[object]]) -> list[int]:
    normalized = [clean(header).lower() for header in headers]
    matches = [index for index, header in enumerate(normalized) if any(preferred in header for preferred in PREFERRED_HEADERS)]
    if matches:
        return matches
    scores = []
    for index in range(len(headers)):
        values = [clean(row[index]) for row in rows if index < len(row)]
        scores.append(sum(len(value) for value in values) / max(len(values), 1))
    return [max(range(len(scores)), key=scores.__getitem__)]


def read_csv(path: Path) -> tuple[list[str], list[list[object]]]:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        records = list(csv.reader(stream))
    if not records:
        return [], []
    return records[0], records[1:]


def read_xlsx(path: Path) -> list[tuple[list[str], list[list[object]]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise SystemExit("Install openpyxl to import .xlsx feedback: python -m pip install openpyxl") from error
    workbook = load_workbook(path, read_only=True, data_only=True)
    collections = []
    for sheet in workbook.worksheets:
        records = list(sheet.iter_rows(values_only=True))
        if records:
            collections.append((list(records[0]), [list(row) for row in records[1:]]))
    return collections


def extract_feedback(headers: list[str], rows: list[list[object]]) -> list[str]:
    if not headers:
        return []
    columns = choose_columns(headers, rows)
    return [clean(row[column]) for row in rows for column in columns if column < len(row)]


def main() -> None:
    sources = sorted(path for path in SOURCE_DIR.iterdir() if path.suffix.lower() in {".csv", ".xlsx"})
    if not sources:
        raise SystemExit(f"Add a CSV or XLSX file to {SOURCE_DIR}")
    source = sources[0]
    if source.suffix.lower() == ".csv":
        headers, rows = read_csv(source)
        collections = [(headers, rows)]
    else:
        collections = read_xlsx(source)
    if not collections:
        raise SystemExit("The spreadsheet is empty.")
    feedbacks: list[str] = []
    seen: set[str] = set()
    for value in (value for headers, rows in collections for value in extract_feedback(headers, rows)):
        identity = value.casefold()
        if len(value) < 4 or identity in PLACEHOLDERS or identity in seen:
            continue
        seen.add(identity)
        feedbacks.append(value)
    OUTPUT.write_text(json.dumps(feedbacks, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Imported {len(feedbacks)} anonymous feedback entries from {source.name}.")


if __name__ == "__main__":
    main()
